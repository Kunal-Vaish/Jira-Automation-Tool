import json
import os
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext
import requests
import csv
from collections import defaultdict
import xlrd
import openpyxl
import tempfile
import traceback
import time


class JiraXrayUploader(tk.Frame):
    def __init__(self, master, controller=None, menu_page_class=None,
                 jira_url="", username="", password=""):
        super().__init__(master)
        self.master = master
        self.controller = controller
        self.menu_page_class = menu_page_class

        # ✅ Credentials passed in from HomePage
        self.jira_url_val = jira_url
        self.username_val = username
        self.password_val = password

        self.grid_rowconfigure(1, weight=1)
        self.grid_columnconfigure(0, weight=1)

        if self.controller and self.menu_page_class:
            top_frame = tk.Frame(self)
            top_frame.grid(row=0, column=0, sticky="nw", padx=10, pady=10)
            self.back_btn = tk.Button(top_frame, text="⬅ Back", width=12, height=1, bg="lightgray",
                                      command=self.go_back)
            self.back_btn.pack(anchor="nw")
        else:
            self.back_btn = None  # standalone mode

        container = tk.Frame(self)
        container.grid(row=1, column=0, padx=20, pady=20, sticky="nsew")

        for i in range(3):
            container.grid_columnconfigure(i, weight=1)

        row = 0
        # ✅ Show logged-in user info
        tk.Label(container, text=f"Logged in as: {self.username_val}", fg="green",
                 font=("Arial", 12, "bold")).grid(row=row, column=0, columnspan=3, pady=5)
        row += 1

        tk.Label(container, text="User Story Number").grid(row=row, column=0, columnspan=3, pady=5)
        row += 1
        self.user_story = tk.Entry(container, width=40)
        self.user_story.grid(row=row, column=0, columnspan=3, pady=5)
        row += 1

        tk.Button(container, text="Select File (.csv/.xls/.xlsx)", command=self.select_file).grid(row=row, column=0, columnspan=3, pady=10)
        row += 1
        self.file_label = tk.Label(container, text="No file selected")
        self.file_label.grid(row=row, column=0, columnspan=3, pady=5)
        row += 1

        self.upload_btn = tk.Button(container, text="Upload Test Cases", command=self.upload)
        self.upload_btn.grid(row=row, column=0, columnspan=3, pady=10)
        row += 1

        self.clear_btn = tk.Button(container, text="Clear Logs", command=self.clear_logs, bg="lightyellow")
        self.clear_btn.grid(row=row, column=0, columnspan=3, pady=5)
        row += 1

        # Paned window for logs
        paned = tk.PanedWindow(container, orient=tk.HORIZONTAL, sashrelief=tk.RAISED)
        paned.grid(row=row, column=0, columnspan=3, sticky="nsew", pady=10)
        container.grid_rowconfigure(row, weight=1)
        container.grid_columnconfigure(0, weight=1)

        left_frame = tk.Frame(paned)
        tk.Label(left_frame, text="General Log", font=("Arial", 10, "bold")).pack(anchor="w")
        self.general_log = scrolledtext.ScrolledText(left_frame, width=60, height=25)
        self.general_log.pack(fill="both", expand=True)
        self.general_log.config(state="disabled")
        paned.add(left_frame)

        right_frame = tk.Frame(paned)
        tk.Label(right_frame, text="Errors", font=("Arial", 10, "bold"), fg="red").pack(anchor="w")
        self.error_log = scrolledtext.ScrolledText(right_frame, width=60, height=25, fg="red")
        self.error_log.pack(fill="both", expand=True)
        self.error_log.config(state="disabled")
        paned.add(right_frame)

        def fix_sash():
            paned.update_idletasks()
            total_width = paned.winfo_width()
            paned.sash_place(0, total_width // 2, 0)
        container.after(100, fix_sash)

        self.file_path = None
        self.temp_csv = None

    def go_back(self):
        if self.controller and self.menu_page_class:
            self.controller.show_frame(self.menu_page_class)
    
    def select_file(self):
        self.file_path = filedialog.askopenfilename(filetypes=[("CSV, XLS, XLSX", "*.csv;*.xls;*.xlsx")])
        self.file_label.config(text=self.file_path or "No file selected")

    def log_message(self, msg):
        self.general_log.config(state="normal")
        self.general_log.insert(tk.END, msg + "\n")
        self.general_log.see(tk.END)
        self.general_log.config(state="disabled")
        self.master.update()

    def log_error(self, msg):
        self.error_log.config(state="normal")
        self.error_log.insert(tk.END, msg + "\n")
        self.error_log.see(tk.END)
        self.error_log.config(state="disabled")
        self.master.update()

    def clear_logs(self):
        self.general_log.config(state="normal")
        self.general_log.delete(1.0, tk.END)
        self.general_log.config(state="disabled")

        self.error_log.config(state="normal")
        self.error_log.delete(1.0, tk.END)
        self.error_log.config(state="disabled")

    def set_inputs_state(self, state):
        widgets = [self.user_story, self.upload_btn]
        for w in widgets:
            w.config(state=state)

    def ensure_csv(self, file_path):
        if file_path.endswith(".csv"):
            return file_path
        elif file_path.endswith(".xls"):
            workbook = xlrd.open_workbook(file_path)
            sheet = workbook.sheet_by_index(0)
            tmp_csv = tempfile.NamedTemporaryFile(delete=False, suffix=".csv", mode="w", newline="", encoding="utf-8")
            writer = csv.writer(tmp_csv)
            for r in range(sheet.nrows):
                writer.writerow(sheet.row_values(r))
            tmp_csv.close()
            self.temp_csv = tmp_csv.name
            return tmp_csv.name
        elif file_path.endswith(".xlsx"):
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheet = wb.active
            tmp_csv = tempfile.NamedTemporaryFile(delete=False, suffix=".csv", mode="w", newline="", encoding="utf-8")
            writer = csv.writer(tmp_csv)
            for row in sheet.iter_rows(values_only=True):
                writer.writerow([cell if cell is not None else "" for cell in row])
            tmp_csv.close()
            self.temp_csv = tmp_csv.name
            return tmp_csv.name
        else:
            raise ValueError("Unsupported file type. Only .csv, .xls, and .xlsx are allowed.")

    def upload(self):
        if not all([self.jira_url_val, self.username_val, self.password_val, self.user_story.get(), self.file_path]):
            messagebox.showerror("Error", "All fields and file selection are required")
            return

        if self.back_btn:
            self.back_btn.config(state="disabled")

        if self.clear_btn:
            self.clear_btn.config(state="disabled")

        self.set_inputs_state("disabled")

        try:
            csv_path = self.ensure_csv(self.file_path)
        except Exception as e:
            messagebox.showerror("Error", f"File conversion failed: {e}")
            self.set_inputs_state("normal")
            return

        auth = (self.username_val, self.password_val)
        headers = {"Content-Type": "application/json"}
        user_story_key = self.user_story.get().strip()
        project_key = user_story_key.split("-")[0]
        url_issue = f"{self.jira_url_val.rstrip('/')}/rest/api/2/issue"

        try:
            # Create Test Plan
            self.log_message("=== Creating Test Plan ===")
            test_plan_payload = {
                "fields": {
                    "project": {"key": project_key},
                    "summary": f"Test Plan for {user_story_key}",
                    "description": f"Created automated Test Plan for {user_story_key} via Kunalv's App",
                    "issuetype": {"name": "Test Plan"},
                    "assignee": {"name": self.username_val},
                }
            }
            resp = requests.post(url_issue, json=test_plan_payload, auth=auth, headers=headers)
            resp.raise_for_status()
            test_plan_key = resp.json()["key"]
            self.log_message(f"✅ Test Plan created: {test_plan_key}")

            # Link Test Plan to User Story
            link_payload = {
                "type": {"name": "Relates"},
                "inwardIssue": {"key": test_plan_key},
                "outwardIssue": {"key": user_story_key},
            }
            requests.post(f"{self.jira_url_val.rstrip('/')}/rest/api/2/issueLink", json=link_payload, auth=auth, headers=headers)

            # Create Test Execution
            self.log_message("=== Creating Test Execution ===")
            test_execution_payload = {
                "fields": {
                    "project": {"key": project_key},
                    "summary": f"Test Execution for {user_story_key}",
                    "description": f"Created automated Test Execution for {user_story_key} via Kunalv's App",
                    "issuetype": {"name": "Test Execution"},
                    "assignee": {"name": self.username_val},
                    "customfield_11368": [test_plan_key]
                }
            }
            resp_exec = requests.post(url_issue, json=test_execution_payload, auth=auth, headers=headers)
            resp_exec.raise_for_status()
            test_execution_key = resp_exec.json()["key"]
            self.log_message(f"✅ Test Execution created: {test_execution_key}")

            # Read CSV and create Test Cases
            self.log_message("=== Creating Test Cases ===")
            tests = defaultdict(list)
            with open(csv_path, newline="", encoding="utf-8") as csvfile:
                reader = csv.DictReader(csvfile)
                for row in reader:
                    tests[row["S.No."].strip()].append(row)

            created_count = 0
            total = len(tests)
            created_tests = []

            for sno, rows in tests.items():
                first = rows[0]
                summary = first.get("Test Name") or f"Test {sno}"
                description = first.get("HLTC Description & Pre-condition") or ""
                priority = first.get("Priority") or "Medium"
                fix_version = first.get("Fix Version/s")
                fix_versions = [{"name": fix_version}] if fix_version else []
                labels = list(filter(None, (first.get("labels") or "").split(",")))
                assignee = self.username_val

                steps = []
                for i, r in enumerate(rows, start=1):
                    action = r.get("LLTC Description") or ""
                    expected = r.get("Expected result") or ""
                    if action:
                        steps.append({
                            "index": i,
                            "fields": {
                                "Action": action,
                                "Data": "",
                                "Expected Result": expected
                            }
                        })

                test_payload = {
                    "fields": {
                        "project": {"key": project_key},
                        "summary": summary,
                        "description": description,
                        "issuetype": {"name": "Test"},
                        "priority": {"name": priority},
                        "fixVersions": fix_versions,
                        "assignee": {"name": assignee},
                        "labels": labels,
                        "customfield_11341": {"id": "12281"},   # Test Type: Manual
                        "customfield_11345": {"steps": steps},  # Steps
                        "customfield_11350": [test_plan_key]   # Link Test Plan
                    }
                }

                try:
                    resp_test = requests.post(url_issue, json=test_payload, auth=auth, headers=headers)
                    resp_test.raise_for_status()
                    test_key = resp_test.json()["key"]
                    created_tests.append(test_key)
                    created_count += 1
                    self.log_message(f"Progress: <{created_count}/{total}> Test Cases have been Uploaded")
                    time.sleep(0.2)
                except requests.exceptions.HTTPError as he:
                    resp = he.response
                    status = resp.status_code
                    try:
                        error_json = resp.json()
                        jira_msg = json.dumps(error_json, ensure_ascii=False, indent=2)
                    except Exception:
                        jira_msg = resp.text
                    self.log_error(f"❌ Failed Test '{summary}': HTTP {status}")
                    self.log_error(f"Jira Response:\n{jira_msg}")
                except Exception as e:
                    self.log_error(f"❌ Failed Test '{summary}': {e}")

            # After creating all tests, bulk add to Test Execution via Xray API
            if created_tests:
                try:
                    self.log_message("=== Adding Tests to Test Execution ===")
                    url_exec_add = f"{self.jira_url_val.rstrip('/')}/rest/raven/1.0/api/testexec/{test_execution_key}/test"
                    test_add_payload = {"add": created_tests}
                    resp_exec_add = requests.post(url_exec_add, json=test_add_payload, auth=auth, headers=headers)
                    resp_exec_add.raise_for_status()
                    self.log_message(f"🔗 Added {len(created_tests)} Tests to Test Execution {test_execution_key}")
                except Exception as e:
                    self.log_error(f"⚠️ Failed to add Tests to Test Execution: {e}")
            else:
                self.log_message("No tests were created successfully; skipping add-to-execution step.")

            self.log_message(f"All done! {created_count}/{total} test cases uploaded under {test_plan_key}")
            messagebox.showinfo("Success", f"Uploaded {created_count}/{total} test cases under {test_plan_key}")

        except Exception as e:
            self.log_error(f"❌ Error: {e}")
            messagebox.showerror("Error", f"Error: {e}")
        finally:
            if self.back_btn:
                self.back_btn.config(state="normal")

            if self.clear_btn:
                self.clear_btn.config(state="normal")

            self.set_inputs_state("normal")

            if self.temp_csv and os.path.exists(self.temp_csv):
                os.remove(self.temp_csv)


if __name__ == "__main__":
    try:
        root = tk.Tk()
        root.title("Jira Test Case Uploader")
        try:
            root.state("zoomed")
        except:
            root.attributes("-zoomed", True)

        # Standalone run: credentials must be manually provided
        # Example placeholders (replace with actual credentials if testing standalone)
        jira_url = "https://your-jira-url"
        username = "your-username"
        password = "your-password"

        app = JiraXrayUploader(root, jira_url=jira_url, username=username, password=password)
        app.pack(fill="both", expand=True)
        root.mainloop()
    except Exception:
        with open("error.log", "w", encoding="utf-8") as f:
            f.write(traceback.format_exc())
        raise